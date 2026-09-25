import os
from openpyxl import load_workbook

p = r'c:\Users\crist\Desktop\cotizador\excel\valores 2027.xlsx'
print('exists:', os.path.exists(p))
if os.path.exists(p):
    wb = load_workbook(p, data_only=True)
    print('sheets:', wb.sheetnames)
    for ws in wb.worksheets:
        print('--- SHEET', ws.title)
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 25), values_only=True):
            print(row)
